"""
Step 19 — Document extraction.

Lets a user upload an actual PDF / DOCX / TXT file instead of pasting raw
text into a textarea. Pasting works, but nobody with a real 40-page policy
document is going to copy-paste it — file upload is what makes the
knowledge-base feature genuinely usable.

Deliberately dependency-light and defensive: a scanned PDF (images, no text
layer) extracts nothing, and an encrypted or corrupt file raises — both are
reported back to the user as a clear message rather than a stack trace or,
worse, a silently empty knowledge base.
"""

import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
# low enough that a short but legitimate document (a one-paragraph policy)
# isn't rejected, high enough to catch a genuinely empty extraction
MIN_USABLE_CHARS = 20


class ExtractionError(Exception):
    """Raised with a user-facing message when a file can't be turned into text."""


def _extract_pdf(data: bytes) -> str:
    try:
        import pdfplumber
    except ImportError:
        raise ExtractionError("PDF support isn't installed. Run: pip install pdfplumber")

    try:
        pages = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        return "\n\n".join(p for p in pages if p.strip())
    except Exception as e:
        raise ExtractionError(f"Couldn't read that PDF ({type(e).__name__}). "
                               "If it's password-protected, remove the password and try again.")


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise ExtractionError("DOCX support isn't installed. Run: pip install python-docx")

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        # tables carry real content in policy/compliance documents, so pull
        # those too rather than silently dropping them
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    except Exception as e:
        raise ExtractionError(f"Couldn't read that Word document ({type(e).__name__}).")


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Couldn't decode that text file — try saving it as UTF-8.")


def extract_text(filename: str, data: bytes) -> str:
    """Returns plain text from an uploaded file, or raises ExtractionError
    with a message intended to be shown directly to the user."""
    if not data:
        raise ExtractionError("That file appears to be empty.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ExtractionError(f"Unsupported file type '{suffix or filename}'. Supported: {supported}")

    if suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".docx":
        text = _extract_docx(data)
    else:
        text = _extract_txt(data)

    text = text.strip()
    if len(text) < MIN_USABLE_CHARS:
        if suffix == ".pdf":
            raise ExtractionError(
                "Almost no text could be extracted from that PDF. If it's a scanned "
                "document (a photo of pages rather than real text), it needs OCR "
                "first — this version doesn't do OCR."
            )
        raise ExtractionError(
            "That file has too little text to build a knowledge base from. "
            "Try a document with more content."
        )
    return text


if __name__ == "__main__":
    # round-trip check on a generated DOCX and TXT; PDF is exercised by the
    # test suite, which builds one on the fly
    import docx as _docx

    doc = _docx.Document()
    doc.add_paragraph("Our refund window is 30 days from delivery.")
    doc.add_paragraph("Refunds are issued to the original payment method.")
    buf = io.BytesIO()
    doc.save(buf)
    print("docx ->", extract_text("policy.docx", buf.getvalue())[:80])

    print("txt  ->", extract_text("notes.txt", b"Support hours are 9am to 6pm on weekdays only.")[:80])

    try:
        extract_text("image.png", b"not a document")
    except ExtractionError as e:
        print("rejects unsupported:", e)

    try:
        extract_text("tiny.txt", b"too short")
    except ExtractionError as e:
        print("rejects too-short:", e)
